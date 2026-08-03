import sys
import os
import csv
import json
import shutil
from pathlib import Path
# Ensure the script's own directory is on sys.path so sibling imports work
# regardless of which directory python is invoked from.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog

SCRIPT_DIR = Path(os.path.dirname(os.path.abspath(__file__)))
ROOT_DIR = SCRIPT_DIR.parent
LX_O_REP_DIR = ROOT_DIR / "LX_O_REP"
OUTPUT_DIRS = [ROOT_DIR / "outputs", ROOT_DIR / "LXDS_O_REP_Output" / "outputs"]

WORKSPACE_SCAN_SOURCES = {
    "Workspace Root": ROOT_DIR,
    "LXDS": SCRIPT_DIR,
    "LX_O_REP": LX_O_REP_DIR,
    "Outputs": OUTPUT_DIRS[0] if OUTPUT_DIRS[0].exists() else ROOT_DIR / "LXDS_O_REP_Output" / "outputs",
}

# --- Dependency Check (no subprocesses / no ghost popups) ---
_dep_root = tk.Tk()
_dep_root.withdraw()
try:
    import xlsxwriter
    import openpyxl
except ImportError as e:
    messagebox.showerror(
        "Missing Dependency",
        f"Missing package: {e.name}\nPlease run: pip install xlsxwriter openpyxl"
    )
    _dep_root.destroy()
    sys.exit(1)
_dep_root.destroy()
# -------------------------------------------------------------

# --------------- Palette ---------------
WHITE_BG    = "#FFFFFF"
LIGHT_GREY  = "#F0F0F0"
DARK_SLATE  = "#2C3E50"
ACCENT_GOLD = "#B5860D"
MID_GREY    = "#BDC3C7"
# ---------------------------------------

# --------------- Data Types for Frame Attributes ---------------
ATTR_DATA_TYPES = [
    "N(Counter)",
    "N(Float)",
    "N(Integer)",
    "AlphaN",
    "String",
    "Custom",
]
# ---------------------------------------------------------------

# =============================================================================
#  Frame Maker Wizard  (3-Step Toplevel)
# =============================================================================

class FrameMakerWizard(tk.Toplevel):
    """
    A 3-step wizard dialog for defining a new LX Frame schema.

    Step 1 — Alias Allocation   : alias length + alias string entry
    Step 2 — Attribute Editor   : dynamic row-based attribute definition
    Step 3 — Review & Generate  : summary + Save & Generate Frame
    """

    def __init__(self, parent):
        super().__init__(parent)
        self.title("Frame Maker Wizard")
        self.geometry("640x520")
        self.resizable(True, True)
        self.configure(bg=WHITE_BG)
        self.grab_set()          # modal

        # --- Wizard state ---
        self.current_step = 1

        # Step-1 data
        self.alias_length_var = tk.StringVar(value="3")
        self.alias_string_var = tk.StringVar()

        # Step-2 data: list of dicts {name_var, type_var, req_var, row_frame}
        self.attr_rows: list[dict] = []

        # --- Header strip ---
        self.header_frame = tk.Frame(self, bg=DARK_SLATE)
        self.header_frame.pack(fill="x")
        self.header_label = tk.Label(
            self.header_frame,
            text="Step 1 — Alias Allocation",
            font=("Segoe UI", 13, "bold"),
            bg=DARK_SLATE, fg="white",
            pady=12
        )
        self.header_label.pack()

        # --- Step indicator ---
        self.step_bar = tk.Frame(self, bg=LIGHT_GREY, height=4)
        self.step_bar.pack(fill="x")
        self.step_canvas = tk.Canvas(self.step_bar, height=4, bg=LIGHT_GREY,
                                     highlightthickness=0)
        self.step_canvas.pack(fill="x")

        # --- Content area (swapped per step) ---
        self.content_frame = tk.Frame(self, bg=WHITE_BG)
        self.content_frame.pack(fill="both", expand=True, padx=20, pady=10)

        # --- Footer (Back / Next / Generate) ---
        self.footer = tk.Frame(self, bg=LIGHT_GREY)
        self.footer.pack(fill="x", side="bottom")

        self.btn_back = tk.Button(
            self.footer, text="← Back",
            command=self._go_back,
            bg=MID_GREY, fg=DARK_SLATE,
            font=("Segoe UI", 10), padx=14, pady=6
        )
        self.btn_back.pack(side="left", padx=12, pady=8)

        self.btn_next = tk.Button(
            self.footer, text="Next →",
            command=self._go_next,
            bg=DARK_SLATE, fg="white",
            font=("Segoe UI", 10, "bold"), padx=14, pady=6
        )
        self.btn_next.pack(side="right", padx=12, pady=8)

        self._render_step(1)

    # ------------------------------------------------------------------
    # Navigation
    # ------------------------------------------------------------------

    def _go_next(self):
        if self.current_step == 1:
            if not self._validate_step1():
                return
            self.current_step = 2
            self._render_step(2)
        elif self.current_step == 2:
            if not self._validate_step2():
                return
            self.current_step = 3
            self._render_step(3)

    def _go_back(self):
        if self.current_step == 2:
            self.current_step = 1
            self._render_step(1)
        elif self.current_step == 3:
            self.current_step = 2
            self._render_step(2)

    # ------------------------------------------------------------------
    # Step Rendering
    # ------------------------------------------------------------------

    def _clear_content(self):
        for w in self.content_frame.winfo_children():
            w.destroy()

    def _render_step(self, step: int):
        self._clear_content()
        self._update_header(step)
        self._update_step_bar(step)

        if step == 1:
            self._build_step1()
            self.btn_back.config(state="disabled")
            self.btn_next.config(text="Next →", command=self._go_next)
        elif step == 2:
            self._build_step2()
            self.btn_back.config(state="normal")
            self.btn_next.config(text="Next →", command=self._go_next)
        elif step == 3:
            self._build_step3()
            self.btn_back.config(state="normal")
            self.btn_next.config(
                text="💾  Save & Generate Frame",
                command=self._generate,
                bg=ACCENT_GOLD
            )

    def _update_header(self, step: int):
        labels = {
            1: "Step 1 — Alias Allocation",
            2: "Step 2 — Attribute Definition",
            3: "Step 3 — Review & Generate",
        }
        self.header_label.config(text=labels[step])

    def _update_step_bar(self, step: int):
        self.step_canvas.update_idletasks()
        w = self.step_canvas.winfo_width() or 640
        self.step_canvas.config(width=w)
        self.step_canvas.delete("all")
        filled = int(w * step / 3)
        self.step_canvas.create_rectangle(0, 0, w, 4, fill="#E0E0E0", outline="")
        self.step_canvas.create_rectangle(0, 0, filled, 4, fill=ACCENT_GOLD, outline="")

    # ------------------------------------------------------------------
    # Step 1 — Alias Allocation
    # ------------------------------------------------------------------

    def _build_step1(self):
        f = self.content_frame

        tk.Label(f, text="Define the Frame Alias for the new schema.",
                 font=("Segoe UI", 10), bg=WHITE_BG, fg="#555555").pack(anchor="w", pady=(8, 16))

        row1 = tk.Frame(f, bg=WHITE_BG)
        row1.pack(fill="x", pady=6)
        tk.Label(row1, text="Frame Alias Length:", width=22, anchor="w",
                 font=("Segoe UI", 10), bg=WHITE_BG).pack(side="left")
        self._entry_alias_len = tk.Entry(row1, textvariable=self.alias_length_var,
                                          width=8, font=("Segoe UI", 10))
        self._entry_alias_len.pack(side="left", padx=4)
        tk.Label(row1, text="(integer, e.g. 3, 4, 5)",
                 font=("Segoe UI", 9), fg="#888888", bg=WHITE_BG).pack(side="left", padx=6)

        row2 = tk.Frame(f, bg=WHITE_BG)
        row2.pack(fill="x", pady=6)
        tk.Label(row2, text="Frame Alias String:", width=22, anchor="w",
                 font=("Segoe UI", 10), bg=WHITE_BG).pack(side="left")
        self._entry_alias_str = tk.Entry(row2, textvariable=self.alias_string_var,
                                          width=20, font=("Segoe UI", 10))
        self._entry_alias_str.pack(side="left", padx=4)
        tk.Label(row2, text="(e.g. IMG, REPRF, HDCREP)",
                 font=("Segoe UI", 9), fg="#888888", bg=WHITE_BG).pack(side="left", padx=6)

        # Live validation hint
        hint_frame = tk.Frame(f, bg="#FFF9EE", bd=1, relief="solid")
        hint_frame.pack(fill="x", pady=(14, 0), ipady=6, ipadx=8)
        tk.Label(hint_frame,
                 text="ℹ  The alias string length must equal the Frame Alias Length value.\n"
                      "   Example: Length=3 → Alias=IMG  |  Length=5 → Alias=REPRF",
                 font=("Segoe UI", 9), bg="#FFF9EE", fg="#7D6000",
                 justify="left").pack(anchor="w", padx=8)

    def _validate_step1(self) -> bool:
        length_str = self.alias_length_var.get().strip()
        alias_str  = self.alias_string_var.get().strip().upper()

        if not length_str.isdigit() or int(length_str) < 1:
            messagebox.showerror("Validation Error",
                                 "Frame Alias Length must be a positive integer.",
                                 parent=self)
            return False

        length = int(length_str)
        if len(alias_str) != length:
            messagebox.showerror("Validation Error",
                                 f"Alias string '{alias_str}' must be exactly {length} characters long.",
                                 parent=self)
            return False

        self.alias_string_var.set(alias_str)
        return True

    # ------------------------------------------------------------------
    # Step 2 — Dynamic Attribute Editor
    # ------------------------------------------------------------------

    def _build_step2(self):
        f = self.content_frame

        tk.Label(f,
                 text="Add parameter attributes for this frame. Each attribute defines a typed field.",
                 font=("Segoe UI", 10), bg=WHITE_BG, fg="#555555",
                 wraplength=580, justify="left").pack(anchor="w", pady=(8, 6))

        # Column headers
        hdr = tk.Frame(f, bg=LIGHT_GREY)
        hdr.pack(fill="x", pady=(0, 2))
        for col_text, col_w in [("Attribute Name", 22), ("Data Type", 16), ("Required", 8), ("", 4)]:
            tk.Label(hdr, text=col_text, width=col_w, anchor="w",
                     font=("Segoe UI", 9, "bold"), bg=LIGHT_GREY, fg=DARK_SLATE,
                     pady=4).pack(side="left", padx=2)

        # Scrollable container for attribute rows
        scroll_container = tk.Frame(f, bg=WHITE_BG, bd=1, relief="sunken")
        scroll_container.pack(fill="both", expand=True)

        canvas = tk.Canvas(scroll_container, bg=WHITE_BG, highlightthickness=0)
        scrollbar = ttk.Scrollbar(scroll_container, orient="vertical", command=canvas.yview)
        self._attr_inner = tk.Frame(canvas, bg=WHITE_BG)

        self._attr_inner.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        canvas.create_window((0, 0), window=self._attr_inner, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Bind mouse-wheel scrolling
        canvas.bind_all("<MouseWheel>",
                        lambda e: canvas.yview_scroll(-1 * (e.delta // 120), "units"))

        self._attr_canvas = canvas

        # Re-render existing rows (on back-navigation)
        existing = list(self.attr_rows)
        self.attr_rows = []
        for row_data in existing:
            self._add_attr_row(
                name=row_data["name_var"].get(),
                dtype=row_data["type_var"].get(),
                req=row_data["req_var"].get()
            )

        if not self.attr_rows:
            self._add_attr_row()

        # Add Attribute button
        btn_row = tk.Frame(f, bg=WHITE_BG)
        btn_row.pack(fill="x", pady=6)
        tk.Button(btn_row, text="＋ Add Attribute",
                  command=self._add_attr_row,
                  bg=LIGHT_GREY, fg=DARK_SLATE,
                  font=("Segoe UI", 9), padx=10, pady=4,
                  relief="groove").pack(side="left")

    def _add_attr_row(self, name="", dtype="String", req=False):
        row_data = {
            "name_var": tk.StringVar(value=name),
            "type_var": tk.StringVar(value=dtype),
            "req_var":  tk.BooleanVar(value=req),
        }

        row_f = tk.Frame(self._attr_inner, bg=WHITE_BG, pady=2)
        row_f.pack(fill="x", padx=4)
        row_data["row_frame"] = row_f

        # Name entry
        name_entry = tk.Entry(row_f, textvariable=row_data["name_var"],
                               width=22, font=("Segoe UI", 10))
        name_entry.pack(side="left", padx=(0, 6))

        # Data type combobox
        type_cb = ttk.Combobox(row_f, textvariable=row_data["type_var"],
                                values=ATTR_DATA_TYPES, width=14,
                                state="readonly", font=("Segoe UI", 10))
        type_cb.pack(side="left", padx=(0, 6))

        # Required checkbox
        req_cb = tk.Checkbutton(row_f, variable=row_data["req_var"],
                                 bg=WHITE_BG, activebackground=WHITE_BG)
        req_cb.pack(side="left", padx=(6, 2))

        # Remove button
        def remove(rd=row_data):
            rd["row_frame"].destroy()
            self.attr_rows.remove(rd)

        tk.Button(row_f, text="✕", command=remove,
                  bg="#E74C3C", fg="white",
                  font=("Segoe UI", 8, "bold"),
                  width=2, pady=1, relief="flat").pack(side="left", padx=4)

        self.attr_rows.append(row_data)
        self._attr_inner.update_idletasks()
        self._attr_canvas.configure(scrollregion=self._attr_canvas.bbox("all"))

    def _validate_step2(self) -> bool:
        if not self.attr_rows:
            messagebox.showerror("Validation Error",
                                 "Add at least one attribute before proceeding.",
                                 parent=self)
            return False
        for i, rd in enumerate(self.attr_rows):
            name = rd["name_var"].get().strip()
            if not name:
                messagebox.showerror("Validation Error",
                                     f"Attribute row {i+1} has an empty name.",
                                     parent=self)
                return False
        return True

    # ------------------------------------------------------------------
    # Step 3 — Review & Generate
    # ------------------------------------------------------------------

    def _build_step3(self):
        f = self.content_frame
        alias  = self.alias_string_var.get()
        length = self.alias_length_var.get()

        tk.Label(f, text="Review your frame specification before generating.",
                 font=("Segoe UI", 10), bg=WHITE_BG, fg="#555555").pack(anchor="w", pady=(8, 12))

        # Summary card
        card = tk.Frame(f, bg="#F8F9FA", bd=1, relief="solid")
        card.pack(fill="x", pady=4, ipady=10)

        tk.Label(card, text=f"  Frame Alias:   {alias}",
                 font=("Segoe UI", 11, "bold"), bg="#F8F9FA", anchor="w").pack(fill="x", padx=12, pady=2)
        tk.Label(card, text=f"  Alias Length:  {length}",
                 font=("Segoe UI", 10), bg="#F8F9FA", anchor="w").pack(fill="x", padx=12, pady=2)
        tk.Label(card, text=f"  Attributes:    {len(self.attr_rows)}",
                 font=("Segoe UI", 10), bg="#F8F9FA", anchor="w").pack(fill="x", padx=12, pady=2)

        # Attribute summary table
        tbl_frame = tk.Frame(f, bg=WHITE_BG)
        tbl_frame.pack(fill="both", expand=True, pady=8)

        # Header
        hdr = tk.Frame(tbl_frame, bg=DARK_SLATE)
        hdr.pack(fill="x")
        for col_text, col_w in [("#", 4), ("Name", 24), ("Type", 16), ("Required", 10)]:
            tk.Label(hdr, text=col_text, width=col_w, anchor="w",
                     font=("Segoe UI", 9, "bold"), bg=DARK_SLATE, fg="white",
                     pady=4).pack(side="left", padx=4)

        # Rows
        for i, rd in enumerate(self.attr_rows):
            bg = WHITE_BG if i % 2 == 0 else "#F5F5F5"
            row_f = tk.Frame(tbl_frame, bg=bg)
            row_f.pack(fill="x")
            tk.Label(row_f, text=str(i+1), width=4, anchor="w",
                     font=("Segoe UI", 9), bg=bg).pack(side="left", padx=4)
            tk.Label(row_f, text=rd["name_var"].get(), width=24, anchor="w",
                     font=("Segoe UI", 9), bg=bg).pack(side="left", padx=4)
            tk.Label(row_f, text=rd["type_var"].get(), width=16, anchor="w",
                     font=("Segoe UI", 9), bg=bg).pack(side="left", padx=4)
            req_text = "✔  Yes" if rd["req_var"].get() else "—  No"
            req_fg   = "#27AE60" if rd["req_var"].get() else "#888888"
            tk.Label(row_f, text=req_text, width=10, anchor="w",
                     font=("Segoe UI", 9), bg=bg, fg=req_fg).pack(side="left", padx=4)

        # Output paths preview
        script_dir   = os.path.dirname(os.path.abspath(__file__))
        json_path    = os.path.join(script_dir, "frame_templates", f"{alias}_frame.json")
        parent_dir   = os.path.abspath(os.path.join(script_dir, ".."))
        py_path      = os.path.join(parent_dir, "LX_O_REP", "frames", f"F_{alias}.py")

        out_card = tk.Frame(f, bg="#EEF8EE", bd=1, relief="solid")
        out_card.pack(fill="x", pady=8, ipady=6)
        tk.Label(out_card, text="Output files:",
                 font=("Segoe UI", 9, "bold"), bg="#EEF8EE", fg="#1A6B1A").pack(anchor="w", padx=10, pady=(4, 0))
        tk.Label(out_card, text=f"  📄  {json_path}",
                 font=("Segoe UI", 8), bg="#EEF8EE", fg="#333333",
                 wraplength=580, justify="left").pack(anchor="w", padx=10)
        tk.Label(out_card, text=f"  🐍  {py_path}",
                 font=("Segoe UI", 8), bg="#EEF8EE", fg="#333333",
                 wraplength=580, justify="left").pack(anchor="w", padx=10, pady=(0, 4))

    # ------------------------------------------------------------------
    # Generation
    # ------------------------------------------------------------------

    def _generate(self):
        alias      = self.alias_string_var.get()
        alias_len  = int(self.alias_length_var.get())
        script_dir = os.path.dirname(os.path.abspath(__file__))
        parent_dir = os.path.abspath(os.path.join(script_dir, ".."))

        # --- Build attribute spec list ---
        attributes = []
        for rd in self.attr_rows:
            attributes.append({
                "name":     rd["name_var"].get().strip(),
                "type":     rd["type_var"].get(),
                "required": rd["req_var"].get(),
            })

        # ---------------------------------------------------------------
        # 1. Write <ALIAS>_frame.json  →  LXDS/frame_templates/
        # ---------------------------------------------------------------
        json_dir  = os.path.join(script_dir, "frame_templates")
        os.makedirs(json_dir, exist_ok=True)
        json_path = os.path.join(json_dir, f"{alias}_frame.json")

        schema = {
            "alias":       alias,
            "alias_length": alias_len,
            "version":     "1.0",
            "attributes":  attributes,
        }

        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(schema, f, indent=4)

        # ---------------------------------------------------------------
        # 2. Write F_<ALIAS>.py  →  LX_O_REP/frames/
        # ---------------------------------------------------------------
        frames_dir = os.path.join(parent_dir, "LX_O_REP", "frames")
        os.makedirs(frames_dir, exist_ok=True)
        py_path = os.path.join(frames_dir, f"F_{alias}.py")

        py_code = self._generate_python_boilerplate(alias, attributes)
        with open(py_path, "w", encoding="utf-8") as f:
            f.write(py_code)

        messagebox.showinfo(
            "Frame Generated",
            f"Frame '{alias}' successfully created!\n\n"
            f"📄 Schema:  {json_path}\n"
            f"🐍 Module:  {py_path}",
            parent=self
        )
        self.destroy()

    # ------------------------------------------------------------------
    # Python Boilerplate Generator
    # ------------------------------------------------------------------

    @staticmethod
    def _generate_python_boilerplate(alias: str, attributes: list) -> str:
        """
        Build a F_<ALIAS>.py skeleton with:
          - LX_BaseFrame inheritance
          - Lock stubs (set_lock_status, set_internal_lock)
          - Per-attribute type validation helpers
          - execute() with field extraction + validation calls
        """
        # Map declared types to Python type hints / validators
        TYPE_MAP = {
            "N(Counter)":  ("int",   "isinstance(v, int) and v >= 0"),
            "N(Float)":    ("float", "isinstance(v, (int, float))"),
            "N(Integer)":  ("int",   "isinstance(v, int)"),
            "AlphaN":      ("str",   "isinstance(v, str) and v.replace('-','').replace('_','').isalnum()"),
            "String":      ("str",   "isinstance(v, str)"),
            "Custom":      ("object","True  # Custom type — implement validation below"),
        }

        lines = []
        lines.append(f'"""')
        lines.append(f"F_{alias}.py — Auto-generated LX Frame Module")
        lines.append(f"Alias   : {alias}")
        lines.append(f"Schema  : {alias}_frame.json")
        lines.append(f"Inherits: LX_BaseFrame")
        lines.append(f'Generated by LXDS Frame Maker Wizard.')
        lines.append(f'"""')
        lines.append("")
        lines.append("import sys")
        lines.append("import os")
        lines.append("sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))")
        lines.append("from LX_BaseFrame import LX_BaseFrame, LXLockState")
        lines.append("")
        lines.append("")
        lines.append(f"class F_{alias}(LX_BaseFrame):")
        lines.append(f'    """Frame handler for alias \'{alias}\' — edit the execute() method to implement logic."""')
        lines.append("")
        lines.append(f"    ALIAS = \"{alias}\"")
        lines.append("")
        lines.append(f"    def __init__(self):")
        lines.append(f"        super().__init__(self.ALIAS)")
        lines.append(f"        # Initialize per-attribute internal locks")
        for attr in attributes:
            lines.append(f"        self.set_internal_lock(\"{attr['name']}\", False)")
        lines.append("")

        # --- Validation helpers per attribute ---
        for attr in attributes:
            name  = attr["name"]
            dtype = attr["type"]
            py_type, check = TYPE_MAP.get(dtype, ("object", "True"))
            req_note = "required" if attr["required"] else "optional"
            lines.append(f"    def _validate_{name.lower()}(self, v) -> bool:")
            lines.append(f"        \"\"\"{name} — declared type: {dtype} ({req_note})\"\"\"")
            lines.append(f"        try:")
            lines.append(f"            result = {check}")
            lines.append(f"            if not result:")
            lines.append(f"                self._log(f\"[WARN] {{self.ALIAS}}: '{name}' failed type check (expected {dtype}). Got: {{v!r}}\")")
            lines.append(f"            return result")
            lines.append(f"        except Exception as exc:")
            lines.append(f"            self._log(f\"[ERROR] {{self.ALIAS}}: '{name}' validation error — {{exc}}\")")
            lines.append(f"            return False")
            lines.append("")

        # --- execute() ---
        lines.append("    def execute(self, row_data: dict) -> bool:")
        lines.append(f'        """')
        lines.append(f"        Main execution entry for '{alias}' frame rows.")
        lines.append(f"        Called by the Orchestrator per matching Record.csv row.")
        lines.append(f'        """')
        lines.append("        self.set_lock_status(LXLockState.BUSY)")
        lines.append("        try:")
        lines.append("            # --- Extract fields from row_data ---")
        for attr in attributes:
            name = attr["name"]
            req  = attr["required"]
            lines.append(f"            {name.lower()} = row_data.get(\"{name}\", \"\")")
            if req:
                lines.append(f"            if not {name.lower()}:")
                lines.append(f"                self._log(f\"[ERROR] Missing required field '{name}' in row: {{row_data}}\")")
                lines.append(f"                self.set_lock_status(LXLockState.ERROR)")
                lines.append(f"                return False")
            lines.append(f"            if not self._validate_{name.lower()}({name.lower()}):")
            lines.append(f"                self._log(f\"[WARN] '{name}' failed validation — skipping.\")")
        lines.append("")
        lines.append("            # --- TODO: Implement frame-specific logic here ---")
        lines.append("")
        lines.append("            self.set_lock_status(LXLockState.FREE)")
        lines.append("            return True")
        lines.append("        except Exception as e:")
        lines.append("            self._log(f\"[ERROR] {alias}.execute() raised: {e}\")")
        lines.append("            self.set_lock_status(LXLockState.ERROR)")
        lines.append("            return False")
        lines.append("")

        # --- _log() helper ---
        lines.append("    def _log(self, msg: str):")
        lines.append('        """Basic console logger — replace with LX_CommCentre hook if needed."""')
        lines.append("        print(msg)")
        lines.append("")

        return "\n".join(lines)


# =============================================================================
#  Editor Tab Components and Workspace Backend
# =============================================================================

class EditableFileTab:
    def __init__(self, app, file_path: Path):
        self.app = app
        self.path = file_path
        self.title = self.path.name
        self.modified = False
        self.frame = tk.Frame(self.app.editor_notebook, bg=WHITE_BG)

    def set_modified(self, value: bool = True):
        self.modified = value
        self._update_tab_title()

    def _update_tab_title(self):
        tab_id = self.app.editor_notebook.select()
        if tab_id:
            current_widget = self.app.editor_notebook.nametowidget(tab_id)
            if current_widget is self.frame:
                title = f"{self.title}{' *' if self.modified else ''}"
                self.app.editor_notebook.tab(tab_id, text=title)

    def save(self):
        raise NotImplementedError("save must be implemented by subclasses")

    def reload(self):
        raise NotImplementedError("reload must be implemented by subclasses")


class TextFileTab(EditableFileTab):
    def __init__(self, app, file_path: Path):
        super().__init__(app, file_path)
        self.text_widget = None
        self._build_ui()
        self.reload()

    def _build_ui(self):
        control_frame = tk.Frame(self.frame, bg=WHITE_BG)
        control_frame.pack(fill="x", padx=10, pady=(10, 6))

        tk.Button(control_frame, text="Save", command=self.save,
                  bg=ACCENT_GOLD, fg="white", font=("Segoe UI", 9, "bold"), padx=12, pady=4).pack(side="left")
        tk.Button(control_frame, text="Reload", command=self.reload,
                  bg=MID_GREY, fg=DARK_SLATE, font=("Segoe UI", 9), padx=12, pady=4).pack(side="left", padx=8)
        tk.Label(control_frame, text=str(self.path), font=("Segoe UI", 8), bg=WHITE_BG, fg="#555555").pack(side="left", padx=14)

        text_frame = tk.Frame(self.frame, bg=WHITE_BG)
        text_frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        self.text_widget = tk.Text(text_frame, wrap="none", font=FONT_BODY,
                                   bg=CARD_WHITE, fg="#222222", undo=True)
        self.text_widget.pack(side="left", fill="both", expand=True)
        self.text_widget.bind("<<Modified>>", self._on_modified)

        vsb = ttk.Scrollbar(text_frame, orient="vertical", command=self.text_widget.yview)
        self.text_widget.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")

        hsb = ttk.Scrollbar(self.frame, orient="horizontal", command=self.text_widget.xview)
        self.text_widget.configure(xscrollcommand=hsb.set)
        hsb.pack(fill="x", padx=10, pady=(0, 10))

    def _on_modified(self, event):
        if self.text_widget.edit_modified():
            self.set_modified(True)
            self.text_widget.edit_modified(False)

    def save(self):
        try:
            content = self.text_widget.get("1.0", tk.END)
            with open(self.path, "w", encoding="utf-8") as f:
                f.write(content.rstrip() + "\n")
            self.set_modified(False)
            self.app.log(f"[OK] Saved text file: {self.path}")
        except Exception as exc:
            messagebox.showerror("Save Failed", f"Unable to save {self.path}: {exc}")

    def reload(self):
        try:
            with open(self.path, "r", encoding="utf-8") as f:
                content = f.read()
            self.text_widget.delete("1.0", tk.END)
            self.text_widget.insert("1.0", content)
            self.set_modified(False)
            self.app.log(f"[OK] Loaded text file: {self.path}")
        except Exception as exc:
            messagebox.showerror("Load Failed", f"Unable to read {self.path}: {exc}")


class CSVFileTab(EditableFileTab):
    def __init__(self, app, file_path: Path):
        super().__init__(app, file_path)
        self.fieldnames = []
        self.rows = []
        self.tree = None
        self._edit_widget = None
        self._build_ui()
        self.reload()

    def _build_ui(self):
        toolbar = tk.Frame(self.frame, bg=WHITE_BG)
        toolbar.pack(fill="x", padx=10, pady=(10, 6))

        tk.Button(toolbar, text="Save CSV", command=self.save,
                  bg=ACCENT_GOLD, fg="white", font=("Segoe UI", 9, "bold"), padx=12, pady=4).pack(side="left")
        tk.Button(toolbar, text="Reload", command=self.reload,
                  bg=MID_GREY, fg=DARK_SLATE, font=("Segoe UI", 9), padx=12, pady=4).pack(side="left", padx=6)
        tk.Button(toolbar, text="Add Row", command=self.add_row,
                  bg=CREME_ALT, fg=DARK_SLATE, font=("Segoe UI", 9), padx=12, pady=4).pack(side="left", padx=6)
        tk.Button(toolbar, text="Delete Row", command=self.delete_selected_row,
                  bg=CREME_ALT, fg=DARK_SLATE, font=("Segoe UI", 9), padx=12, pady=4).pack(side="left", padx=6)
        tk.Button(toolbar, text="Toggle LX", command=self.toggle_lx_column,
                  bg=CREME_ALT, fg=DARK_SLATE, font=("Segoe UI", 9), padx=12, pady=4).pack(side="left", padx=6)
        tk.Button(toolbar, text="Recalibrate Index", command=self.recalibrate_indexes,
                  bg=CREME_ALT, fg=DARK_SLATE, font=("Segoe UI", 9), padx=12, pady=4).pack(side="left", padx=6)
        tk.Label(toolbar, text=str(self.path), font=("Segoe UI", 8), bg=WHITE_BG, fg="#555555").pack(side="left", padx=14)

        grid_frame = tk.Frame(self.frame, bg=CARD_WHITE, bd=1, relief="solid")
        grid_frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        self.tree = ttk.Treeview(grid_frame, show="headings", selectmode="browse")
        self.tree.pack(side="left", fill="both", expand=True)
        self.tree.bind("<Double-1>", self._on_double_click)

        vsb = ttk.Scrollbar(grid_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(self.frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        vsb.pack(side="right", fill="y")
        hsb.pack(fill="x", padx=10, pady=(0, 10))

    def reload(self):
        try:
            with open(self.path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                self.fieldnames = reader.fieldnames or []
                self.rows = [dict(row) for row in reader]

            if not self.fieldnames:
                messagebox.showwarning("CSV Load", f"{self.path} has no header row.")
                self.fieldnames = []
                self.rows = []

            self._render_table()
            self.set_modified(False)
            self.app.log(f"[OK] Loaded CSV file: {self.path}")
        except Exception as exc:
            messagebox.showerror("CSV Load Failed", f"Unable to read {self.path}: {exc}")

    def save(self):
        if not self.fieldnames:
            messagebox.showerror("Save Failed", f"CSV file {self.path} has no headers.")
            return

        try:
            self._commit_active_edit()
            with open(self.path, "w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=self.fieldnames)
                writer.writeheader()
                for row in self.rows:
                    writer.writerow(row)

            self.set_modified(False)
            self.app.log(f"[OK] Saved CSV file: {self.path}")
        except Exception as exc:
            messagebox.showerror("Save Failed", f"Unable to save {self.path}: {exc}")

    def _render_table(self):
        self.tree.delete(*self.tree.get_children())
        self.tree["columns"] = self.fieldnames
        for field in self.fieldnames:
            self.tree.heading(field, text=field)
            self.tree.column(field, width=140, anchor="w")

        for index, row in enumerate(self.rows):
            values = [row.get(field, "") for field in self.fieldnames]
            tag = "even" if index % 2 == 0 else "odd"
            self.tree.insert("", "end", iid=str(index), values=values, tags=(tag,))
            self.tree.tag_configure("even", background=CARD_WHITE)
            self.tree.tag_configure("odd", background=CREME_BG)

    def _on_double_click(self, event):
        region = self.tree.identify("region", event.x, event.y)
        if region != "cell":
            return
        item = self.tree.identify_row(event.y)
        column = self.tree.identify_column(event.x)
        if not item or not column:
            return
        self._start_cell_edit(item, column)

    def _start_cell_edit(self, item, column_id):
        self._commit_active_edit()
        column_index = int(column_id.replace("#", "")) - 1
        field = self.fieldnames[column_index]
        bbox = self.tree.bbox(item, column_id)
        if not bbox:
            return
        x, y, width, height = bbox
        value = self.tree.set(item, field)

        self._edit_widget = tk.Entry(self.tree)
        self._edit_widget.insert(0, value)
        self._edit_widget.place(x=x, y=y, width=width, height=height)
        self._edit_widget.focus_set()
        self._edit_widget.bind("<Return>", lambda e: self._commit_cell(item, field))
        self._edit_widget.bind("<FocusOut>", lambda e: self._commit_cell(item, field))

    def _commit_cell(self, item, field):
        if not self._edit_widget:
            return
        new_value = self._edit_widget.get()
        self._edit_widget.destroy()
        self._edit_widget = None

        row_index = int(item)
        self.rows[row_index][field] = new_value
        self.tree.set(item, field, new_value)
        self.set_modified(True)

    def _commit_active_edit(self):
        if self._edit_widget:
            self._edit_widget.destroy()
            self._edit_widget = None

    def add_row(self):
        row = {field: "" for field in self.fieldnames}
        self.rows.append(row)
        self._render_table()
        self.set_modified(True)

    def delete_selected_row(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Delete Row", "Select a row to delete.")
            return
        index = int(selected[0])
        self.rows.pop(index)
        self._render_table()
        self.set_modified(True)

    def toggle_lx_column(self):
        if "LX" not in self.fieldnames:
            messagebox.showwarning("Toggle LX", "No LX column found in this CSV.")
            return
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Toggle LX", "Select a row to toggle LX.")
            return
        idx = int(selected[0])
        current = str(self.rows[idx].get("LX", "")).strip().upper()
        self.rows[idx]["LX"] = "" if current == "X" else "X"
        self._render_table()
        self.set_modified(True)

    def recalibrate_indexes(self):
        if "Index" not in self.fieldnames:
            messagebox.showwarning("Recalibrate Index", "No Index column available.")
            return
        for i, row in enumerate(self.rows, start=1):
            row["Index"] = str(i)
        self._render_table()
        self.set_modified(True)
        self.app.log(f"[OK] Recalibrated Index column for {self.path}")


class AgentCLI(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app.root)
        self.app = app
        self.title("Agent CLI")
        self.geometry("720x420")
        self.configure(bg=WHITE_BG)
        self._build_ui()

    def _build_ui(self):
        tk.Label(self, text="Agent Command Interface", font=("Segoe UI", 13, "bold"),
                 bg=WHITE_BG, fg=DARK_SLATE).pack(anchor="w", padx=14, pady=(14, 8))

        self.output = tk.Text(self, height=16, bg=CARD_WHITE, fg="#222222", font=FONT_BODY, state="disabled")
        self.output.pack(fill="both", expand=True, padx=14, pady=(0, 8))

        command_frame = tk.Frame(self, bg=WHITE_BG)
        command_frame.pack(fill="x", padx=14, pady=(0, 14))

        self.command_entry = tk.Entry(command_frame, font=FONT_BODY)
        self.command_entry.pack(side="left", fill="x", expand=True, padx=(0, 12))
        self.command_entry.bind("<Return>", lambda e: self.execute_command())

        tk.Button(command_frame, text="Run", command=self.execute_command,
                  bg=ACCENT_GOLD, fg="white", font=("Segoe UI", 9, "bold"), padx=14, pady=6).pack(side="right")

        self._print("Type 'help' for available commands.")

    def _print(self, msg: str):
        self.output.config(state="normal")
        self.output.insert(tk.END, msg + "\n")
        self.output.see(tk.END)
        self.output.config(state="disabled")

    def execute_command(self):
        command = self.command_entry.get().strip()
        self.command_entry.delete(0, tk.END)
        if not command:
            return
        self._print(f"> {command}")
        parts = command.split(maxsplit=1)
        action = parts[0].lower()
        argument = parts[1] if len(parts) > 1 else ""

        if action == "help":
            self._print("Commands: help, list, open <path>, load <csv>, save, newframe, refresh, close")
        elif action == "list":
            self._print("Open files:")
            for path in self.app.open_tabs:
                self._print(f" - {path}")
        elif action == "open":
            if not argument:
                self._print("Usage: open <path>")
            else:
                self.app.open_file(Path(argument))
        elif action == "load":
            if not argument:
                self._print("Usage: load <csv>")
            else:
                self.app.open_csv_file(Path(argument))
        elif action == "save":
            self.app.save_all()
            self._print("Saved all open files.")
        elif action == "newframe":
            self.app.frame_maker()
            self._print("Frame Maker launched.")
        elif action == "refresh":
            self.app.populate_file_tree()
            self._print("Workspace Explorer refreshed.")
        elif action == "close":
            self.destroy()
        else:
            self._print(f"Unknown command: {action}")


# =============================================================================
#  Main Application
# =============================================================================

class LXDS_App:
    def __init__(self, root):
        self.root = root
        self.root.title("LX Development Suite (LXDS)")
        self.root.geometry("1000x700")
        self.root.configure(bg=WHITE_BG)

        self.is_admin = False
        self.open_tabs: dict[str, EditableFileTab] = {}
        self.file_tree = None
        self.editor_notebook = None
        self.empty_editor_label = None
        self.log_area = None

        self.create_login_screen()

    def create_login_screen(self):
        self.clear_frame()

        login_frame = tk.Frame(self.root, bg=WHITE_BG)
        login_frame.place(relx=0.5, rely=0.5, anchor="center")

        tk.Label(login_frame, text="LXDS Developer Mode",
                 font=("Segoe UI", 16, "bold"), bg=WHITE_BG, fg=DARK_SLATE).pack(pady=20)

        tk.Label(login_frame, text="Auth Key (Password):",
                 font=("Segoe UI", 10), bg=WHITE_BG).pack()
        self.entry_pass = tk.Entry(login_frame, show="*", width=30, font=("Segoe UI", 10))
        self.entry_pass.pack(pady=5)
        self.entry_pass.bind("<Return>", lambda e: self.check_login())

        tk.Button(login_frame, text="Unlock",
                  command=self.check_login,
                  bg=DARK_SLATE, fg="white",
                  font=("Segoe UI", 10, "bold"),
                  padx=18, pady=6, relief="flat").pack(pady=20)

    def check_login(self):
        pwd = self.entry_pass.get()
        self.is_admin = pwd in ("alexvsakha", "EXTREMIS")
        if pwd == "EXTREMIS":
            self.launch_extremis()
        self.create_main_dashboard()

    def clear_frame(self):
        for widget in self.root.winfo_children():
            widget.destroy()

    def create_main_dashboard(self):
        self.clear_frame()

        menubar = tk.Menu(self.root)
        filemenu = tk.Menu(menubar, tearoff=0)
        filemenu.add_command(label="Open Contract",  command=self.load_contract)
        filemenu.add_command(label="Load SrcOT",     command=self.load_srcot)
        filemenu.add_command(label="Export Scheme",  command=self.export_scheme)
        filemenu.add_separator()
        filemenu.add_command(label="Exit",           command=self.root.quit)
        menubar.add_cascade(label="File", menu=filemenu)

        toolsmenu = tk.Menu(menubar, tearoff=0)
        toolsmenu.add_command(label="Frame Maker",         command=self.frame_maker)
        toolsmenu.add_command(label="Counter Recalibrator",command=self.counter_recalibrator)
        toolsmenu.add_command(label="Block Preset Inserter",command=self.open_block_inserter)
        toolsmenu.add_command(label="Bit Splitter",        command=self.launch_extremis)
        menubar.add_cascade(label="Tools", menu=toolsmenu)

        contractsmenu = tk.Menu(menubar, tearoff=0)
        contractsmenu.add_command(label="Validate Matrix",          command=self.validate_matrix)
        contractsmenu.add_command(label="Sync Schema",              command=self.sync_schema)
        contractsmenu.add_command(label="Verify Double-Link Integrity", command=self.verify_double_link_integrity)
        menubar.add_cascade(label="Contracts", menu=contractsmenu)

        cli_menu = tk.Menu(menubar, tearoff=0)
        cli_menu.add_command(label="Open Dedicated Agentic Command Interface", command=self.open_agent_cli)
        menubar.add_cascade(label="Agent CLI", menu=cli_menu)

        about_menu = tk.Menu(menubar, tearoff=0)
        about_menu.add_command(label="AXC Lab Licensing & Version Info",
                                command=lambda: messagebox.showinfo("About", "AXC Lab - LX Programs Suite v2.0"))
        menubar.add_cascade(label="About", menu=about_menu)

        self.root.config(menu=menubar)
        self.root.bind_all("<Control-s>", lambda e: self.save_all())
        self.root.bind_all("<Control-S>", lambda e: self.save_all())

        style = ttk.Style()
        style.configure("TFrame", background=WHITE_BG)
        style.configure("TNotebook", background=LIGHT_GREY)
        style.configure("Treeview", font=("Segoe UI", 9), rowheight=22)
        style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"))

        paned = tk.PanedWindow(self.root, orient=tk.HORIZONTAL, bg=LIGHT_GREY)
        paned.pack(fill=tk.BOTH, expand=True)

        left_nav = tk.Frame(paned, bg=DARK_SLATE, width=280)
        paned.add(left_nav)

        tk.Label(left_nav, text="LXDS Workspace Explorer",
                 font=("Segoe UI", 12, "bold"), bg=DARK_SLATE, fg="white").pack(pady=14)
        tk.Frame(left_nav, bg=ACCENT_GOLD, height=2).pack(fill="x", padx=10, pady=(0, 10))

        self.file_tree = ttk.Treeview(left_nav, show="tree")
        self.file_tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 8))
        self.file_tree.bind("<Double-1>", self.on_tree_double_click)
        self.file_tree.bind("<Button-3>", self.on_tree_right_click)

        self.file_tree_menu = tk.Menu(self.root, tearoff=0)
        self.file_tree_menu.add_command(label="New Contract", command=self.create_new_contract)
        self.file_tree_menu.add_command(label="New Frame", command=self.create_frame_and_refresh)
        self.file_tree_menu.add_command(label="Load SrcOT", command=self.load_srcot)
        self.file_tree_menu.add_separator()
        self.file_tree_menu.add_command(label="Rename", command=self.rename_selected_file)
        self.file_tree_menu.add_command(label="Delete", command=self.delete_selected_file)

        quick_actions = tk.Frame(left_nav, bg=DARK_SLATE)
        quick_actions.pack(fill=tk.X, padx=10, pady=10)
        for label, cmd in [
            ("Open Contract", self.load_contract),
            ("Load SrcOT", self.load_srcot),
            ("Import CSV", self.import_multi_srcot),
            ("Save All", self.save_all),
        ]:
            tk.Button(quick_actions, text=label, command=cmd,
                      bg="#3D5166", fg="white", font=("Segoe UI", 9), relief="flat",
                      activebackground=ACCENT_GOLD, activeforeground="white",
                      pady=6).pack(fill=tk.X, pady=3)

        right_frame = tk.Frame(paned, bg=LIGHT_GREY)
        paned.add(right_frame, stretch="always")

        toolbar = tk.Frame(right_frame, bg=WHITE_BG, pady=10)
        toolbar.pack(fill=tk.X, padx=10)
        for label, cmd in [
            ("Load SrcOT", self.load_srcot),
            ("Load Contract", self.load_contract),
            ("Import Multi-SrcOT", self.import_multi_srcot),
            ("Save All", self.save_all),
        ]:
            tk.Button(toolbar, text=label, command=cmd,
                      bg=DARK_SLATE, fg="white", font=("Segoe UI", 9, "bold"), padx=14, pady=6,
                      relief="flat").pack(side=tk.LEFT, padx=4)

        content_notebook = ttk.Notebook(right_frame)
        content_notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 10))

        editor_tab = ttk.Frame(content_notebook)
        arch_tab = ttk.Frame(content_notebook)
        log_tab = ttk.Frame(content_notebook)

        content_notebook.add(editor_tab, text="Editor")
        content_notebook.add(arch_tab, text="Architecture")
        content_notebook.add(log_tab, text="Logs")

        self.editor_notebook = ttk.Notebook(editor_tab)
        self.editor_notebook.pack(fill=tk.BOTH, expand=True, padx=6, pady=6)

        self.empty_editor_label = tk.Label(self.editor_notebook,
                                           text="Open a file from the workspace explorer to begin editing.",
                                           font=("Segoe UI", 11), bg=WHITE_BG, fg="#666666")
        self.empty_editor_label.pack(fill=tk.BOTH, expand=True, padx=20, pady=20)
        self.editor_notebook.add(self.empty_editor_label, text="Welcome")

        self.arch_text = tk.Text(arch_tab, wrap="word", bg=CARD_WHITE, fg="#222222",
                                 font=("Segoe UI", 10), padx=12, pady=12)
        self.arch_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self.arch_text.insert("1.0", self._get_architecture_spec_text())
        self.arch_text.config(state="disabled")

        self.log_area = tk.Text(log_tab, height=10, bg=CARD_WHITE, fg="#222222",
                                 font=("Consolas", 10), bd=1, relief="solid")
        self.log_area.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        self.populate_file_tree()
        self.log("LXDS loaded: workspace explorer and editor are ready.")

    def _get_architecture_spec_text(self) -> str:
        return (
            "LX Architecture Specification\n"
            "===============================\n\n"
            "Orchestrator -> Neuron N_* -> Organ O_* -> Frame F_* -> LXComm Boundary & Lock Engine\n\n"
            "1. Orchestrator\n"
            "   - Central engine that reads Source of Truth rows and routes execution to frame handlers.\n"
            "   - Validates record rows, resolves destination paths, and writes report outputs via xlsxwriter.\n\n"
            "2. Neuron N_*\n"
            "   - Conceptual layer for distributed row processing and stateful execution.\n"
            "   - Each record row may be treated as a neuron input, with mirror references and sequence pointers.\n\n"
            "3. Organ O_*\n"
            "   - High-level path and report management orchestrated by LX_CommCentre.\n"
            "   - Dynamic path placeholders P0..P9 and PCB_ID resolution are resolved before output creation.\n\n"
            "4. Frame F_*\n"
            "   - Frame modules implement domain-specific logic for row Types such as IMG, MIR, HDC, VAL, TXT, STR.\n"
            "   - All frame modules inherit LX_BaseFrame and expose execute(row_data: dict) -> bool.\n\n"
            "5. LXComm Boundary & Lock Engine\n"
            "   - LX_BaseFrame provides lock state management using FREE, BUSY, SKIP, ERROR states.\n"
            "   - Frame execution respects lock status and provides deterministic handover between row operations.\n\n"
            "Runtime Flow\n"
            "-----------\n"
            "- LXDS builds and manages frame schemas, contract files, and Source of Truth CSVs.\n"
            "- O_REP loads Record.csv, applies PCB_ID and path mapping, then dispatches rows to the Orchestrator.\n"
            "- The Orchestrator evaluates rows by Type and passes them into the matching frame handler.\n"
            "- Frame modules log state changes and update lock status for error containment or successful completion.\n\n"
            "Path Mapping\n"
            "------------\n"
            "- PCB_ID is generated from Project, Version, and Board No.\n"
            "- Placeholders P0..P9 are resolved by LX_CommCentre and can be defined in O_REP UI.\n"
            "- Example mapping: P0 -> C:/Reports, PCB_ID -> MYPROJECT_1.0_01, path string: P0/PCB_ID/Output.xlsx\n\n"
            "Frame Routing\n"
            "-------------\n"
            "- Types register to frame handlers using lookup tables in LX_Orchestrator.\n"
            "- Common aliases: IMG, MIR, HDC, VAL, TXT, STR.\n"
            "- Pointer logic: Mirror fields such as M1, M2 reference existing Index values to create chained or mirrored outputs.\n\n"
            "Lock States\n"
            "-----------\n"
            "- FREE: Frame handler is ready for new execution.\n"
            "- BUSY: The frame is actively processing a row.\n"
            "- SKIP: The row is intentionally skipped.\n"
            "- ERROR: An execution failure occurred and should halt the generated report path.\n\n"
            "AI Agent Integration\n"
            "---------------------\n"
            "- Agents can hook into LXDS by editing frame schemas, writing new F_* modules, or generating contract aliases.\n"
            "- LXDS exposes the workspace via the explorer and supports loading and saving CSV/JSON source-of-truth files directly.\n"
            "- External orchestrators can be connected to O_REP by sharing valid Record.csv/Record_Contract.csv pairs and generating frame modules that conform to LX_BaseFrame.\n"
        )

    def log(self, msg: str):
        if self.log_area and isinstance(self.log_area, tk.Text):
            self.log_area.config(state="normal")
            self.log_area.insert(tk.END, msg + "\n")
            self.log_area.see(tk.END)
            self.log_area.config(state="disabled")
        else:
            print(msg)

    def _resolve_known_file(self, name: str) -> Path | None:
        candidates = [ROOT_DIR / name, SCRIPT_DIR / name, LX_O_REP_DIR / name]
        for p in candidates:
            if p.is_file():
                return p
        return None

    def populate_file_tree(self):
        self.file_tree.delete(*self.file_tree.get_children())
        for label, source in WORKSPACE_SCAN_SOURCES.items():
            if not source.exists():
                continue
            node = self.file_tree.insert("", "end", iid=str(source), text=f"{label}: {source.name}")
            self._populate_tree_children(node, source)

    def _populate_tree_children(self, parent, path: Path):
        try:
            entries = sorted(path.iterdir(), key=lambda x: (x.is_file(), x.name.lower()))
        except PermissionError:
            return
        for child in entries:
            if child.name.startswith("."):
                continue
            if child.is_dir():
                node = self.file_tree.insert(parent, "end", iid=str(child), text=child.name)
                self._populate_tree_children(node, child)
            else:
                self.file_tree.insert(parent, "end", iid=str(child), text=child.name)

    def on_tree_double_click(self, event):
        item = self.file_tree.selection()
        if not item:
            return
        path = Path(item[0])
        if path.is_file():
            self.open_file(path)

    def on_tree_right_click(self, event):
        item = self.file_tree.identify_row(event.y)
        if item:
            self.file_tree.selection_set(item)
        self.file_tree_menu.tk_popup(event.x_root, event.y_root)

    def create_new_contract(self):
        path = filedialog.asksaveasfilename(
            title="Create New Contract CSV",
            defaultextension=".csv",
            initialfile="Record_Contract.csv",
            filetypes=[("CSV Files", "*.csv")]
        )
        if not path:
            return
        fieldnames = ["Alias", "Description", "Type", "Required"]
        try:
            with open(path, "w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
            self.log(f"[OK] Created contract file: {path}")
            self.open_file(Path(path))
            self.populate_file_tree()
        except Exception as exc:
            messagebox.showerror("Create Contract", f"Unable to create contract file: {exc}")

    def create_frame_and_refresh(self):
        wizard = FrameMakerWizard(self.root)
        self.root.wait_window(wizard)
        self.populate_file_tree()

    def open_file(self, file_path: Path):
        file_path = file_path.resolve()
        if str(file_path) in self.open_tabs:
            self.editor_notebook.select(self.open_tabs[str(file_path)].frame)
            return

        if self.empty_editor_label and self.empty_editor_label.winfo_exists():
            self.editor_notebook.forget(self.empty_editor_label)

        if file_path.suffix.lower() == ".csv":
            tab = CSVFileTab(self, file_path)
        else:
            tab = TextFileTab(self, file_path)

        self.open_tabs[str(file_path)] = tab
        self.editor_notebook.add(tab.frame, text=tab.title)
        self.editor_notebook.select(tab.frame)

    def open_csv_file(self, file_path: Path):
        if not file_path.exists():
            messagebox.showwarning("Open CSV", f"File not found: {file_path}")
            return
        self.open_file(file_path)

    def save_all(self):
        for path, tab in self.open_tabs.items():
            try:
                tab.save()
            except Exception as exc:
                self.log(f"[ERROR] Failed saving {path}: {exc}")
        self.log("[OK] Save All complete.")

    def load_srcot(self):
        srcot = self._resolve_known_file("Record.csv")
        if srcot:
            self.open_file(srcot)
        else:
            path = filedialog.askopenfilename(title="Select Record.csv", filetypes=[("CSV Files", "*.csv")])
            if path:
                self.open_file(Path(path))

    def load_contract(self):
        contract = self._resolve_known_file("Record_Contract.csv")
        if contract:
            self.open_file(contract)
        else:
            path = filedialog.askopenfilename(title="Select Record_Contract.csv", filetypes=[("CSV Files", "*.csv")])
            if path:
                self.open_file(Path(path))

    def import_multi_srcot(self):
        paths = filedialog.askopenfilenames(title="Import Source-of-Truth CSVs",
                                            filetypes=[("CSV Files", "*.csv"), ("All Files", "*")])
        for filename in paths:
            self.open_file(Path(filename))

    def export_scheme(self):
        selected = self.editor_notebook.select()
        if not selected:
            messagebox.showwarning("Export Scheme", "Open a file first.")
            return
        widget = self.editor_notebook.nametowidget(selected)
        for tab in self.open_tabs.values():
            if tab.frame is widget:
                dest = filedialog.asksaveasfilename(title="Export Scheme", defaultextension=tab.path.suffix,
                                                    initialfile=tab.path.name, filetypes=[("All Files", "*.*")])
                if dest:
                    shutil.copy(tab.path, dest)
                    self.log(f"[OK] Exported scheme {tab.path} → {dest}")
                return
        messagebox.showwarning("Export Scheme", "Unable to determine the active file.")

    def rename_selected_file(self):
        item = self.file_tree.selection()
        if not item:
            return
        path = Path(item[0])
        if not path.exists():
            return
        new_name = simpledialog.askstring("Rename File", "Enter new file name:", initialvalue=path.name)
        if not new_name:
            return
        target = path.with_name(new_name)
        try:
            path.rename(target)
            self.populate_file_tree()
            self.log(f"[OK] Renamed {path.name} → {target.name}")
        except Exception as exc:
            messagebox.showerror("Rename File", f"Failed to rename: {exc}")

    def delete_selected_file(self):
        item = self.file_tree.selection()
        if not item:
            return
        path = Path(item[0])
        if not path.exists():
            return
        if messagebox.askyesno("Delete File", f"Delete {path.name}? This cannot be undone."):
            try:
                if path.is_file():
                    path.unlink()
                else:
                    shutil.rmtree(path)
                self.populate_file_tree()
                self.log(f"[OK] Deleted {path}")
            except Exception as exc:
                messagebox.showerror("Delete File", f"Unable to delete: {exc}")

    def counter_recalibrator(self):
        selected = self.editor_notebook.select()
        if not selected:
            messagebox.showwarning("Counter Recalibrator", "Open a CSV file first.")
            return
        widget = self.editor_notebook.nametowidget(selected)
        for tab in self.open_tabs.values():
            if tab.frame is widget and isinstance(tab, CSVFileTab):
                tab.recalibrate_indexes()
                return
        messagebox.showwarning("Counter Recalibrator", "Select a CSV tab to recalibrate.")

    def open_block_inserter(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Macro Block Inserter")
        dialog.geometry("640x420")
        dialog.configure(bg=WHITE_BG)

        tk.Label(dialog, text="Insert a reusable block into the active file.",
                 font=("Segoe UI", 12, "bold"), bg=WHITE_BG, fg=DARK_SLATE).pack(anchor="w", padx=14, pady=12)

        block_text = (
            "# LXDS Block Preset\n"
            "# Inserted block begins here. Modify as needed.\n"
            "def block_preset():\n"
            "    print('LXDS macro block inserted')\n"
            "\n"
        )

        txt = tk.Text(dialog, wrap="word", bg=CARD_WHITE, fg="#222222",
                      font=("Consolas", 10), height=12)
        txt.insert("1.0", block_text)
        txt.pack(fill=tk.BOTH, expand=True, padx=14, pady=(0, 10))

        def insert_block():
            selected = self.editor_notebook.select()
            if not selected:
                messagebox.showwarning("Macro Block Inserter", "Open a text file first.")
                return
            widget = self.editor_notebook.nametowidget(selected)
            for tab in self.open_tabs.values():
                if tab.frame is widget and isinstance(tab, TextFileTab):
                    cursor = tab.text_widget.index(tk.INSERT)
                    tab.text_widget.insert(cursor, txt.get("1.0", tk.END))
                    tab.set_modified(True)
                    self.log(f"[OK] Inserted macro block into {tab.path}")
                    dialog.destroy()
                    return
            messagebox.showwarning("Macro Block Inserter", "Active tab must be a text file.")

        tk.Button(dialog, text="Insert Block", command=insert_block,
                  bg=ACCENT_GOLD, fg="white", font=("Segoe UI", 10, "bold"), padx=16, pady=8).pack(pady=10)

    def validate_matrix(self):
        srcot = self._resolve_known_file("Record.csv")
        if srcot:
            self._run_csv_validation(srcot)
        else:
            messagebox.showwarning("Validate Matrix", "Record.csv not found. Load Record.csv first.")

    def _run_csv_validation(self, csv_path: Path):
        try:
            with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                rows = [dict(row) for row in reader]
        except Exception as exc:
            messagebox.showerror("Validate Matrix", f"Unable to read {csv_path}: {exc}")
            return

        contract_path = self._resolve_known_file("Record_Contract.csv")
        valid_types = set()
        if contract_path:
            with open(contract_path, "r", encoding="utf-8-sig", newline="") as f:
                valid_types = {row.get("Alias", "").strip() for row in csv.DictReader(f)}

        self.log("")
        self.log("╔═════════════════════════════════════════════════════════════════════")
        self.log("  MATRIX VALIDATION — Record.csv")
        self.log("╚═════════════════════════════════════════════════════════════════════")

        valid_indices = {row.get("Index", "").strip() for row in rows if row.get("Index", "").strip()}
        issues = 0
        for row_num, row in enumerate(rows, start=1):
            row_issues = []
            if not row.get("Type", "").strip():
                row_issues.append("Type is required")
            if not row.get("D_File_Name", "").strip():
                row_issues.append("D_File_Name is required")
            row_type = row.get("Type", "").strip()
            if valid_types and row_type and row_type not in valid_types:
                row_issues.append(f"Unknown Type alias '{row_type}'")
            mirror = row.get("M", "").strip()
            if mirror and re.match(r"^M\d+$", mirror, re.IGNORECASE):
                target = mirror[1:]
                if target not in valid_indices:
                    row_issues.append(f"Mirror pointer {mirror} references missing Index {target}")
            for coord_field in ["S_Start_Cell_Range", "S_End_Cell_Range", "D_Start_Cell_Range", "D_End_Cell_Range"]:
                coord = row.get(coord_field, "").strip()
                if coord and not re.match(r"^[A-Z]{1,3}[0-9]+$", coord):
                    row_issues.append(f"Invalid coordinate '{coord}' in {coord_field}")
            if row_issues:
                issues += len(row_issues)
                self.log(f"[FAIL] Row {row_num} (Index={row.get('Index', '?')})")
                for issue in row_issues:
                    self.log(f"       - {issue}")
            else:
                self.log(f"[PASS] Row {row_num} (Index={row.get('Index', '?')})")

        if issues == 0:
            self.log("[OK] Record.csv validation passed with no issues.")
        else:
            self.log(f"[WARN] Record.csv validation found {issues} issue(s).")

    def sync_schema(self):
        contract_path = self._resolve_known_file("Record_Contract.csv")
        if not contract_path:
            messagebox.showwarning("Sync Schema", "Record_Contract.csv not found. Load the contract file first.")
            return

        aliases = []
        with open(contract_path, "r", encoding="utf-8-sig", newline="") as f:
            aliases = [row.get("Alias", "").strip() for row in csv.DictReader(f) if row.get("Alias", "").strip()]

        frame_dir = SCRIPT_DIR / "frame_templates"
        frames_dir = LX_O_REP_DIR / "frames"
        frame_dir.mkdir(exist_ok=True)
        frames_dir.mkdir(exist_ok=True)

        generated = []
        for alias in sorted(set(aliases)):
            json_path = frame_dir / f"{alias}_frame.json"
            py_path = frames_dir / f"F_{alias}.py"
            if not json_path.exists():
                json_path.write_text(json.dumps({"alias": alias, "alias_length": len(alias), "version": "1.0", "attributes": []}, indent=4), encoding="utf-8")
                generated.append(json_path.name)
            if not py_path.exists():
                py_path.write_text(self._generate_frame_skeleton(alias), encoding="utf-8")
                generated.append(py_path.name)

        if generated:
            self.log(f"[SYNC] Generated missing schema artifacts: {', '.join(generated)}")
            messagebox.showinfo("Sync Schema", "Schema sync completed and missing artifacts were generated.")
        else:
            self.log("[SYNC] All contract aliases have matching frame artifacts.")
            messagebox.showinfo("Sync Schema", "Schema sync completed. No missing artifacts found.")
        self.populate_file_tree()

    def _generate_frame_skeleton(self, alias: str) -> str:
        return (
            f"'''Auto-generated skeleton for frame alias {alias}'''\n"
            "import os\n"
            "import sys\n"
            "sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))\n"
            "from LX_BaseFrame import LX_BaseFrame, LXLockState\n\n"
            f"class F_{alias}(LX_BaseFrame):\n"
            "    def __init__(self):\n"
            f"        super().__init__(self.ALIAS)\n\n"
            "    def execute(self, row_data: dict) -> bool:\n"
            "        self.set_lock_status(LXLockState.BUSY)\n"
            "        try:\n"
            "            self._log(f'Executing frame {self.ALIAS}')\n"
            "            self.set_lock_status(LXLockState.FREE)\n"
            "            return True\n"
            "        except Exception as exc:\n"
            "            self._log(f'Error: {exc}')\n"
            "            self.set_lock_status(LXLockState.ERROR)\n"
            "            return False\n\n"
            "    def _log(self, msg: str):\n"
            "        print(msg)\n"
        )

    def verify_double_link_integrity(self):
        srcot = self._resolve_known_file("Record.csv")
        if not srcot:
            messagebox.showwarning("Verify Double-Link Integrity", "Record.csv not found. Load it first.")
            return

        try:
            with open(srcot, "r", encoding="utf-8-sig", newline="") as f:
                rows = [dict(row) for row in csv.DictReader(f)]
        except Exception as exc:
            messagebox.showerror("Verify Integrity", f"Unable to read Record.csv: {exc}")
            return

        valid_indices = {row.get("Index", "").strip() for row in rows if row.get("Index", "").strip()}
        broken = []
        for row in rows:
            mirror = row.get("M", "").strip()
            if mirror and re.match(r"^M\d+$", mirror, re.IGNORECASE):
                target = mirror[1:]
                if target not in valid_indices:
                    broken.append((row.get("Index", "?"), mirror))

        if broken:
            for index, mirror in broken:
                self.log(f"[ERROR] Row Index={index} has broken mirror pointer {mirror}")
            messagebox.showwarning("Verify Integrity", f"Found {len(broken)} broken mirror link(s). See logs.")
        else:
            self.log("[OK] Double-link integrity verified for Record.csv.")
            messagebox.showinfo("Verify Integrity", "No broken mirror links detected.")

    def open_agent_cli(self):
        AgentCLI(self)


def main():
    root = tk.Tk()
    LXDS_App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
